"""Join master.csv with the tag files and write the committee spreadsheet.

Output: build/APT_Bulletin_Content_Index.xlsx
"""
import csv, json, glob, collections
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).parent
B = HERE / "build"
TAGS = [("DOC", "TC-Doc"), ("MAT", "TCM (Materials)"), ("PETC", "PETC"), ("CODES", "Codes & Standards"),
        ("MH", "TC-MH (Modern)"), ("SP", "TC-SP (Sustainable)"), ("DIS", "Disaster (PETC subgroup)")]
DISC = [("ENG", "Engineering"), ("ARCH", "Architecture"), ("CONS", "Conservation")]
CONTENT = {"Article", "Practice Point", "Article (type not indexed; may be a review)"}

rows = list(csv.DictReader(open(B / "master.csv")))
# IDs are renumbered whenever build_master.py changes, so tags are joined on (title, year, authors) taken from
# the classify inputs they were made from, not on ID.
key = lambda d: (d["title"], str(d["year"]), d["authors"])
by_key = {}
for f in sorted(glob.glob(str(B / "classify/out_*.json"))):
    inp = {d["id"]: d for d in json.load(open(f.replace("out_", "in_")))}
    for t in json.load(open(f)): by_key[key(inp[t["id"]])] = t
tags = {r["id"]: by_key[key(r)] for r in rows if key(r) in by_key}
# Batches differed on letters/editorials with no clear subject (some left discipline empty, some guessed CONS);
# make them consistent: low-confidence front-of-book items get no discipline.
NONCONTENT = {"Letter", "Editorial", "Guest editor note", "Other", "Note", "Final Touch", "Computer Forum"}
for r in rows:
    t = tags.get(r["id"])
    if t and r["type"] in NONCONTENT and t.get("conf") == "low": t["disc"] = []
missing = [r["id"] for r in rows if r["type"] != "Book reviews (section)" and r["id"] not in tags]

HEAD = Font(bold=True, color="FFFFFF"); FILL = PatternFill("solid", fgColor="1F4E79")

def sheet(ws, header, data, widths):
    ws.append(header)
    for c in ws[1]: c.font, c.fill, c.alignment = HEAD, FILL, Alignment(wrap_text=True, vertical="top")
    for d in data: ws.append(d)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(header))}{max(2, len(data) + 1)}"
    for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w

wb = Workbook()
header = ["ID", "Year", "Vol", "Issue", "Page", "Type", "PP #", "Title", "Authors", "Book author(s) (reviews)", "Topics (APT index headings)"] \
         + [n for _, n in TAGS] + [n for _, n in DISC] + ["Tag confidence", "Reviewed", "Committee notes", "Source", "DOI / link", "Practice Point PDF"]
data = []
for r in rows:
    t = tags.get(r["id"], {})
    link = f"https://doi.org/{r['doi']}" if r["doi"] else ""
    data.append([r["id"], int(r["year"]) if r["year"] else None, int(r["volume"]) if r["volume"] else None, r["issue"],
                 int(r["page_start"]) if (r["page_start"] or "").isdigit() else None, r["type"],
                 int(r["pp_number"]) if r["pp_number"].isdigit() else None, r["title"], r["authors"], r["book_authors"], r["topics"]]
                + [k in t.get("tags", []) for k, _ in TAGS] + [k in t.get("disc", []) for k, _ in DISC]
                + [t.get("conf", ""), False, "", r["source"], link, r["pp_pdf"]])
# Becca only wants articles and Practice Points on the main sheet; book reviews and front/back-of-book
# items (letters, editorials, notes) go on their own sheets. "Book reviews (section)" headers are dropped.
REVIEWS = {"Book review", "Book review or minor item (index only)"}
drop = lambda d, i: d[:i] + d[i + 1:]
def item_sheet(ws, rows_, keep_book_authors):
    h = header if keep_book_authors else drop(header, 9)
    d = rows_ if keep_book_authors else [drop(x, 9) for x in rows_]
    w = [10, 6, 5, 6, 5, 16, 5, 60, 32] + ([22] if keep_book_authors else []) + [40] + [9] * 10 + [9, 9, 30, 12, 26, 26]
    sheet(ws, h, d, w)
    first = h.index(TAGS[0][1]) + 1; last = first + len(TAGS) + len(DISC) - 1
    dv = DataValidation(type="list", formula1='"TRUE,FALSE"', allow_blank=True); ws.add_data_validation(dv)
    dv.add(f"{get_column_letter(first)}2:{get_column_letter(last)}{len(d) + 1}")
    rc = get_column_letter(h.index("Reviewed") + 1); dv.add(f"{rc}2:{rc}{len(d) + 1}")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            if c.column in (8, 9, 10, 11): c.alignment = Alignment(wrap_text=False)
        for c in (row[-2], row[-1]):
            if c.value: c.hyperlink, c.style = c.value, "Hyperlink"
main = [d for d in data if d[5] in CONTENT]
item_sheet(wb.active, main, False); wb.active.title = "Articles & Practice Points"

# Practice Points
pp = sorted((d for d in data if d[6]), key=lambda d: d[6])
ws2 = wb.create_sheet("Practice Points")
sheet(ws2, ["PP #", "Year", "Vol", "Issue", "Title", "Authors", "Topics"] + [n for _, n in TAGS] + [n for _, n in DISC] + ["PDF"],
      [[d[6], d[1], d[2], d[3], d[7], d[8], d[10]] + d[11:11 + len(TAGS) + len(DISC)] + [d[-1]] for d in pp],
      [5, 6, 5, 6, 60, 36, 40] + [9] * 10 + [30])

# Summary: content items (articles + Practice Points) per decade x tag
ws3 = wb.create_sheet("Summary")
content = [d for d in data if d[5] in CONTENT and d[1]]
dec = lambda y: f"{y // 10 * 10}s"
decades = sorted({dec(d[1]) for d in content})
cols = TAGS + DISC
ws3.append(["Articles + Practice Points per decade (from suggested tags; one item can count in several columns)"]); ws3["A1"].font = Font(bold=True)
ws3.append(["Decade", "Items"] + [n for _, n in cols] + ["No committee tag"])
for c in ws3[2]: c.font, c.fill = HEAD, FILL
for dd in decades + ["All"]:
    sub = [d for d in content if dd == "All" or dec(d[1]) == dd]
    flags = [sum(1 for d in sub if d[11 + i]) for i in range(len(cols))]
    none = sum(1 for d in sub if not any(d[11:11 + len(TAGS)]))
    ws3.append([dd, len(sub)] + flags + [none])
ws3.append([])
ws3.append(["Practice Points only"]); ws3.cell(ws3.max_row, 1).font = Font(bold=True)
ws3.append(["", len(pp)] + [sum(1 for d in pp if d[11 + i]) for i in range(len(cols))] + [sum(1 for d in pp if not any(d[11:11 + len(TAGS)]))])
ws3.column_dimensions["A"].width = 14
for i in range(2, len(cols) + 4): ws3.column_dimensions[get_column_letter(i)].width = 13

item_sheet(wb.create_sheet("Book reviews"), [d for d in data if d[5] in REVIEWS], True)
item_sheet(wb.create_sheet("Letters, editorials, other"),
           [d for d in data if d[5] not in CONTENT | REVIEWS | {"Book reviews (section)"}], False)

# Sources and gaps
ws4 = wb.create_sheet("About")
about = open(HERE / "README.md").read().splitlines()
for line in about: ws4.append([line])
ws4.column_dimensions["A"].width = 140

# Rejects
ws5 = wb.create_sheet("Index rows not parsed")
rej = list(csv.DictReader(open(B / "rejects.csv")))
sheet(ws5, ["Index heading", "Entry", "Ref", "Reason"], [[r["category"], r["entry"], r["ref"], r["reason"]] for r in rej], [30, 70, 30, 14])

out = B / "APT_Bulletin_Content_Index.xlsx"
wb.save(out)
print(out, len(main), "articles+PPs;", len(pp), "Practice Points;", len(missing), "untagged:", missing[:10])
c = collections.Counter(d[5] for d in data); print(c)
